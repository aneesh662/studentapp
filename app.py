import os, uuid
from pathlib import Path
from io import BytesIO
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, abort
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import or_
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_ROOT = Path(os.environ.get("UPLOAD_FOLDER", str(BASE_DIR / "uploads")))
PHOTO_DIR, DOCUMENT_DIR = UPLOAD_ROOT / "photos", UPLOAD_ROOT / "documents"
PHOTO_DIR.mkdir(parents=True, exist_ok=True)
DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)
PHOTO_EXTS, PDF_EXTS = {"jpg", "jpeg", "png", "webp"}, {"pdf"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-secret-key")
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES
database_url = os.environ.get("DATABASE_URL", "").strip()
if database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql://", 1)
if not database_url:
    database_url = "sqlite:///" + str(BASE_DIR / "students_local.db")
    print("WARNING: DATABASE_URL not set; using local SQLite for development.")
app.config.update(SQLALCHEMY_DATABASE_URI=database_url, SQLALCHEMY_TRACK_MODIFICATIONS=False,
                  SQLALCHEMY_ENGINE_OPTIONS={"pool_pre_ping": True})
db = SQLAlchemy(app)

class Student(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    phone = db.Column(db.String(40))
    place = db.Column(db.String(120))
    course = db.Column(db.String(120))
    email = db.Column(db.String(160))
    joined_date = db.Column(db.Date)
    photo_filename = db.Column(db.String(255))
    document_filename = db.Column(db.String(255))
    document_original_name = db.Column(db.String(255))
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)

def ext(name):
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""

def save_upload(storage, allowed, folder, label):
    if not storage or not storage.filename:
        return None
    original = secure_filename(storage.filename)
    if not original or ext(original) not in allowed:
        raise ValueError(f"Invalid {label} file type.")
    stored = f"{uuid.uuid4().hex}.{ext(original)}"
    storage.save(folder / stored)
    return stored

def remove_file(folder, filename):
    if filename:
        try:
            (folder / Path(filename).name).unlink(missing_ok=True)
        except OSError:
            app.logger.warning("Could not remove upload %s", filename)

def filtered_students():
    query, search = Student.query, request.args.get("q", "").strip()
    if search:
        like = f"%{search}%"
        query = query.filter(or_(Student.name.ilike(like), Student.phone.ilike(like),
            Student.place.ilike(like), Student.course.ilike(like), Student.email.ilike(like)))
    return query.order_by(Student.id.desc()), search

def form_fields():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Student name is required.", "error")
        return None
    joined = request.form.get("joined_date", "").strip()
    try:
        joined_date = datetime.strptime(joined, "%Y-%m-%d").date() if joined else None
    except ValueError:
        flash("Please enter a valid joining date.", "error")
        return None
    return dict(name=name, phone=request.form.get("phone", "").strip() or None,
        place=request.form.get("place", "").strip() or None,
        course=request.form.get("course", "").strip() or None,
        email=request.form.get("email", "").strip() or None, joined_date=joined_date)

def apply_uploads(student):
    photo, document = request.files.get("photo"), request.files.get("document")
    new_photo = new_pdf = None
    try:
        if photo and photo.filename:
            new_photo = save_upload(photo, PHOTO_EXTS, PHOTO_DIR, "photo")
        if document and document.filename:
            new_pdf = save_upload(document, PDF_EXTS, DOCUMENT_DIR, "PDF document")
    except ValueError as exc:
        remove_file(PHOTO_DIR, new_photo)
        remove_file(DOCUMENT_DIR, new_pdf)
        flash(f"{exc} Photos: JPG/PNG/WEBP. Documents: PDF only.", "error")
        return False
    if new_photo:
        remove_file(PHOTO_DIR, student.photo_filename)
        student.photo_filename = new_photo
    if new_pdf:
        remove_file(DOCUMENT_DIR, student.document_filename)
        student.document_filename = new_pdf
        student.document_original_name = secure_filename(document.filename)
    return True

@app.get("/")
def index():
    query, search = filtered_students()
    return render_template("index.html", students=query.all(), search=search, total=Student.query.count())

@app.post("/students/add")
def add_student():
    fields = form_fields()
    if fields is None:
        return redirect(url_for("index"))
    student = Student(**fields)
    try:
        db.session.add(student)
        db.session.flush()
        if not apply_uploads(student):
            db.session.rollback()
            return redirect(url_for("index"))
        db.session.commit()
        flash("Student added successfully.", "success")
    except Exception:
        db.session.rollback()
        app.logger.exception("Add student failed")
        flash("Could not save student. Check server logs.", "error")
    return redirect(url_for("index"))

@app.route("/students/<int:student_id>/edit", methods=["GET", "POST"])
def edit_student(student_id):
    student = Student.query.get_or_404(student_id)
    if request.method == "POST":
        fields = form_fields()
        if fields is not None:
            for key, value in fields.items():
                setattr(student, key, value)
            if not apply_uploads(student):
                db.session.rollback()
                return redirect(url_for("edit_student", student_id=student.id))
            try:
                db.session.commit()
                flash("Student details updated.", "success")
                return redirect(url_for("index"))
            except Exception:
                db.session.rollback()
                app.logger.exception("Edit student failed")
                flash("Could not update student. Check server logs.", "error")
    return render_template("edit.html", student=student)

@app.post("/students/<int:student_id>/delete")
def delete_student(student_id):
    student = Student.query.get_or_404(student_id)
    photo, pdf = student.photo_filename, student.document_filename
    try:
        db.session.delete(student)
        db.session.commit()
        remove_file(PHOTO_DIR, photo)
        remove_file(DOCUMENT_DIR, pdf)
        flash("Student record deleted.", "success")
    except Exception:
        db.session.rollback()
        app.logger.exception("Delete student failed")
        flash("Could not delete student. Check server logs.", "error")
    return redirect(url_for("index", q=request.form.get("q", "").strip()))

@app.get("/students/<int:student_id>/photo")
def student_photo(student_id):
    student = Student.query.get_or_404(student_id)
    if not student.photo_filename:
        abort(404)
    path = PHOTO_DIR / Path(student.photo_filename).name
    if not path.is_file():
        abort(404)
    return send_file(path, conditional=True)

@app.get("/students/<int:student_id>/document")
def view_document(student_id):
    student = Student.query.get_or_404(student_id)
    if not student.document_filename:
        abort(404)
    path = DOCUMENT_DIR / Path(student.document_filename).name
    if not path.is_file():
        abort(404)
    return send_file(path, mimetype="application/pdf", as_attachment=False,
        download_name=student.document_original_name or "student-document.pdf")

@app.get("/students/export.xlsx")
def export_students():
    query, search = filtered_students()
    wb, sheet = Workbook(), None
    sheet = wb.active
    sheet.title = "Students"
    sheet.append(["ID","Name","Phone","Place","Course","Email","Joined Date","Photo Uploaded","PDF Document"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        cell.alignment = Alignment(horizontal="center")
    for s in query.all():
        sheet.append([s.id,s.name,s.phone or "",s.place or "",s.course or "",s.email or "",
            s.joined_date.isoformat() if s.joined_date else "",
            "Yes" if s.photo_filename else "No", s.document_original_name or ("Yes" if s.document_filename else "No")])
    for col in sheet.columns:
        length = max((len(str(c.value or "")) for c in col), default=10)
        sheet.column_dimensions[get_column_letter(col[0].column)].width = min(length + 3, 40)
    sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="students_filtered.xlsx" if search else "students.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/health")
def health():
    try:
        db.session.execute(db.text("SELECT 1"))
        return {"status":"ok","database":"connected"}, 200
    except Exception:
        app.logger.exception("Health check failed")
        return {"status":"error","database":"unavailable"}, 503

with app.app_context():
    db.create_all()

if __name__ == "__main__":
    app.run(debug=True)
